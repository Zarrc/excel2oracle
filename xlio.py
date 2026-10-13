# -*- coding: utf-8 -*-
"""Excel(.xlsx) 读取与表结构推导 —— 不依赖 Oracle，可独立单元测试。"""
from __future__ import annotations

import re
from datetime import datetime

from openpyxl import load_workbook

MAX_IDENTIFIER = 30      # Oracle 12.1 的标识符上限（12.2+ 为 128，取保守值）
VARCHAR_MAX = 4000       # VARCHAR2 CHAR 语义上限
INSERT_BATCH = 1000      # executemany 批大小
COMMENT_MAX = 300        # COMMENT 文本截断长度（中文约 3 字节/字，留足余量）

_IDENT_RE = re.compile(r"[^0-9A-Za-z_$#]")


def sanitize_identifier(raw, fallback: str) -> str:
    """清洗成合法 Oracle 非引号标识符：大写、去非法字符、不以数字开头、≤30 字符。

    中文表头会被清成空串 → 用 fallback（COL_N）兜底，原始表头由 COMMENT 承载。
    """
    s = _IDENT_RE.sub("_", str(raw if raw is not None else "").strip())
    s = re.sub(r"_+", "_", s).strip("_")
    if not s:
        s = fallback
    if s[0].isdigit():
        s = "C_" + s
    return s[:MAX_IDENTIFIER].upper()


def sanitize_table_name(raw, default: str) -> str:
    """表名清洗，允许 SCHEMA.TABLE 两段；空值 → default。"""
    raw = str(raw if raw is not None else "").strip()
    if not raw:
        return default.upper()
    parts = [p for p in raw.split(".") if p.strip()]
    if len(parts) > 2:
        raise ValueError("表名最多一段 OWNER.TABLE，收到多余点号")
    return ".".join(sanitize_identifier(p, "T") for p in parts)


def build_columns(headers) -> list[tuple[str, str]]:
    """表头 → [(Oracle列名, 原始表头)]：清洗 + 去重（重名加 _2/_3 后缀）。"""
    used: set[str] = set()
    out: list[tuple[str, str]] = []
    for i, h in enumerate(headers, 1):
        base = sanitize_identifier(h, f"COL_{i}")
        name, k = base, 1
        while name in used:
            k += 1
            suffix = f"_{k}"
            # 先给后缀留位再截断（审查 #2：先截断会砍掉后缀导致死循环）
            name = f"{base[:MAX_IDENTIFIER - len(suffix)]}{suffix}"
        used.add(name)
        header_text = "" if h is None else str(h).strip()
        out.append((name, header_text))
    return out


def list_sheets(path: str) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


class SheetData:
    def __init__(self, name: str, columns, rows):
        self.name = name        # sheet 名
        self.columns = columns  # [(Oracle列名, 原始表头)]
        self.rows = rows        # [[v1, v2, ...]] 已规范化（空串→None，bool→TRUE/FALSE）


def read_sheet(path: str, sheet_name: str | None = None) -> SheetData:
    """读一个 sheet：首行为表头；尾部全空表头列、全空行丢弃。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            raise ValueError(f"sheet「{ws.title}」是空的")
        ncols = 0
        for i, h in enumerate(header or [], 1):
            if h is not None and str(h).strip():
                ncols = i
        if ncols == 0:
            raise ValueError(f"sheet「{ws.title}」第一行（表头）全为空")
        columns = build_columns(list(header[:ncols]))
        rows = []
        for r in it:
            if r is None:
                continue
            vals = list(r[:ncols]) + [None] * (ncols - len(r))
            if any(v is not None and str(v) != "" for v in vals):
                rows.append([_normalize(v) for v in vals])
        return SheetData(ws.title, columns, rows)
    finally:
        wb.close()


def _normalize(v):
    """单元格值 → Python 值：bool→TRUE/FALSE 文本，空串→None，其余保型。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (datetime, int, float)):
        return v
    s = str(v)
    return s if s != "" else None


def _to_text(v) -> str:
    """值 → 文本（VARCHAR2 列用）：整数值不带 .0，其余 str()。"""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def infer_column_type(values) -> str:
    """按该列非空样本推导 Oracle 类型：全数值→NUMBER，全日期→DATE/TIMESTAMP，否则 VARCHAR2。"""
    nonnull = [v for v in values if v is not None]
    if not nonnull:
        return f"VARCHAR2({VARCHAR_MAX} CHAR)"
    if all(isinstance(v, datetime) for v in nonnull):
        has_time = any(v.hour or v.minute or v.second or v.microsecond for v in nonnull)
        return "TIMESTAMP" if has_time else "DATE"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in nonnull):
        return "NUMBER"
    texts = [v if isinstance(v, str) else _to_text(v) for v in nonnull]
    return f"VARCHAR2({min(max(len(t) for t in texts), VARCHAR_MAX)} CHAR)"


def infer_types(rows, ncols: int, enabled: bool) -> list[str]:
    """逐列类型；enabled=False 时全 VARCHAR2（staging 稳妥默认）。"""
    if not enabled:
        return [f"VARCHAR2({VARCHAR_MAX} CHAR)"] * ncols
    return [infer_column_type([r[i] for r in rows]) for i in range(ncols)]


def _coerce(v, oracle_type: str):
    """按目标列类型定型，保证 executemany 每列类型同质（None 除外）。"""
    if v is None:
        return None
    if oracle_type.startswith("VARCHAR2"):
        return v if isinstance(v, str) else _to_text(v)
    if oracle_type == "NUMBER":
        if isinstance(v, (int, float)):
            return v
        s = str(v)
        try:
            return int(s)
        except ValueError:
            try:
                return float(s)
            except ValueError:
                return v  # 推断下不应出现；原样交给数据库报错，如实暴露问题
    return v  # DATE / TIMESTAMP：类型推导保证全为 datetime


def coerce_rows(rows, types):
    return [[_coerce(v, t) for v, t in zip(r, types)] for r in rows]


def sql_str(s: str) -> str:
    """SQL 单引号字面量（COMMENT 用）：先截断再转义，保证不把引号对劈开（审查 #4）。"""
    return "'" + str(s)[:COMMENT_MAX].replace("'", "''") + "'"


def make_ddl(table: str, columns, types, source_desc: str) -> list[str]:
    """生成 CREATE TABLE + COMMENT 语句列表（表名/列名均已在上游清洗过）。"""
    cols_sql = ",\n  ".join(f"{name} {t}" for (name, _), t in zip(columns, types))
    stmts = [f"CREATE TABLE {table} (\n  {cols_sql}\n)",
             f"COMMENT ON TABLE {table} IS {sql_str(source_desc)}"]
    for name, header in columns:
        if header:
            stmts.append(f"COMMENT ON COLUMN {table}.{name} IS {sql_str(header)}")
    return stmts


def make_insert_sql(table: str, columns) -> str:
    collist = ", ".join(name for name, _ in columns)
    binds = ", ".join(f":{i + 1}" for i in range(len(columns)))
    return f"INSERT INTO {table} ({collist}) VALUES ({binds})"
