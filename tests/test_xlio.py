# -*- coding: utf-8 -*-
"""xlio 单元测试 —— 不依赖 Oracle 实例。"""
import sys
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # 让 `pytest` 直跑可导入 xlio
import xlio


# ---------- 标识符清洗 ----------
def test_sanitize_basic():
    assert xlio.sanitize_identifier("Order ID", "COL_1") == "ORDER_ID"
    assert xlio.sanitize_identifier("1ABC", "COL_1") == "C_1ABC"
    assert xlio.sanitize_identifier("数量", "COL_1") == "COL_1"      # 中文 → fallback
    assert xlio.sanitize_identifier(None, "COL_2") == "COL_2"
    assert xlio.sanitize_identifier("a" * 50, "X") == "A" * 30       # 截断到 30
    assert xlio.sanitize_identifier("a-b_c", "X") == "A_B_C"


def test_table_name():
    assert xlio.sanitize_table_name("dev.stg_x", "T") == "DEV.STG_X"
    assert xlio.sanitize_table_name("", "STG_DATA") == "STG_DATA"
    assert xlio.sanitize_table_name("  stg order ", "T") == "STG_ORDER"
    with pytest.raises(ValueError):
        xlio.sanitize_table_name("a.b.c", "T")   # 不允许两个以上点号


def test_dedupe_columns():
    cols = xlio.build_columns(["A", "A", "数量"])
    assert [c for c, _ in cols] == ["A", "A_2", "COL_3"]
    assert cols[2][1] == "数量"   # 原始表头保留，供 COMMENT 使用


def test_dedupe_long_duplicate_headers_terminate():
    """回归（审查 HIGH #2）：30 字符重名表头不得死循环，后缀要留位。"""
    long_hdr = "X" * 30
    cols = xlio.build_columns([long_hdr, long_hdr, long_hdr])
    names = [c for c, _ in cols]
    assert len(set(names)) == 3                       # 能走到这里 = 没死循环
    assert all(len(n) <= 30 for n in names)


def test_sql_str_truncates_then_escapes():
    """回归（审查 #4）：截断不劈开引号对 —— 单引号落在切口也应成对。"""
    s = "a" * (xlio.COMMENT_MAX - 1) + "'"
    out = xlio.sql_str(s)
    assert out.startswith("'") and out.endswith("'")
    inner = out[1:-1]
    assert "'" not in inner.replace("''", "")         # 无落单引号


# ---------- 类型推导 ----------
def test_infer_number():
    assert xlio.infer_column_type([1, 2.5, None]) == "NUMBER"


def test_infer_date_vs_timestamp():
    assert xlio.infer_column_type([datetime(2026, 1, 1), None]) == "DATE"
    assert xlio.infer_column_type([datetime(2026, 1, 1, 9, 30)]) == "TIMESTAMP"


def test_infer_varchar_and_empty():
    assert xlio.infer_column_type(["abc", 1, None]) == "VARCHAR2(3 CHAR)"
    assert xlio.infer_column_type([None, None]) == "VARCHAR2(4000 CHAR)"


def test_infer_types_disabled_is_all_varchar():
    ts = xlio.infer_types([[1, "x"]], 2, enabled=False)
    assert ts == ["VARCHAR2(4000 CHAR)", "VARCHAR2(4000 CHAR)"]


def test_coerce_varchar_collapses_float():
    assert xlio._coerce(2.0, "VARCHAR2(10 CHAR)") == "2"
    assert xlio._coerce(2.5, "VARCHAR2(10 CHAR)") == "2.5"
    assert xlio._coerce(None, "NUMBER") is None
    assert xlio._coerce(datetime(2026, 1, 1), "DATE") == datetime(2026, 1, 1)


# ---------- DDL / SQL 生成 ----------
def test_make_ddl_and_insert_sql():
    cols = [("A", "编号"), ("B", "")]
    types = ["NUMBER", "VARCHAR2(10 CHAR)"]
    ddl = xlio.make_ddl("STG_T", cols, types, "src '测试'")
    assert any(s.startswith("CREATE TABLE STG_T") and "A NUMBER" in s
               and "B VARCHAR2(10 CHAR)" in s for s in ddl)
    assert any("COMMENT ON COLUMN STG_T.A IS '编号'" in s for s in ddl)
    assert not any("COMMENT ON COLUMN STG_T.B" in s for s in ddl)   # 空表头不加注释
    assert "src ''测试''" in ddl[1]                                   # 单引号转义
    assert xlio.make_insert_sql("T", cols) == "INSERT INTO T (A, B) VALUES (:1, :2)"


# ---------- Excel 读取（现场造 xlsx） ----------
@pytest.fixture
def sample_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["Order ID", "数量", "", "数量"])   # 第 3 列表头空但有数据；第 4 列与前列重名
    ws.append([1, "十个", None, "x"])
    ws.append([2.0, "", True, None])              # 空串→None，bool→"TRUE"
    p = tmp_path / "s.xlsx"
    wb.save(p)
    return p


def test_read_sheet(sample_xlsx):
    data = xlio.read_sheet(str(sample_xlsx), "数据")
    assert data.name == "数据"
    # 中文表头 fallback 带位置编号（COL_2 / COL_4），天然不冲突；原始表头保留在第二元组
    assert [c for c, _ in data.columns] == ["ORDER_ID", "COL_2", "COL_3", "COL_4"]
    assert data.columns[1][1] == "数量"
    assert len(data.rows) == 2
    r1, r2 = data.rows
    assert r1 == [1, "十个", None, "x"]
    assert r2 == [2.0, None, "TRUE", None]


def test_trailing_empty_header_columns_dropped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["A", "", ""])
    ws.append(["x", 1, 2])
    ws.append([None, None, None])                  # 整行空 → 跳过
    p = tmp_path / "t.xlsx"
    wb.save(p)
    data = xlio.read_sheet(str(p))
    assert [c for c, _ in data.columns] == ["A"]
    assert data.rows == [["x"]]


def test_empty_sheet_raises(tmp_path):
    wb = Workbook()
    wb.active.title = "空"
    p = tmp_path / "e.xlsx"
    wb.save(p)
    with pytest.raises(ValueError):
        xlio.read_sheet(str(p), "空")


def test_list_sheets(sample_xlsx):
    assert xlio.list_sheets(str(sample_xlsx)) == ["数据"]
